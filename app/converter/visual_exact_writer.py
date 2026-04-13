from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _set_section(section, width_pt: float, height_pt: float) -> None:
    section.page_width = Pt(width_pt)
    section.page_height = Pt(height_pt)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)


def _render_page_image(page: fitz.Page, image_path: Path, dpi: int) -> None:
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=matrix, alpha=False)
    pix.save(str(image_path))


def write_docx_visual_exact(pdf_path: str, output_path: str, temp_dir: Path, dpi: int = 220) -> int:
    temp_dir.mkdir(parents=True, exist_ok=True)
    document = Document()

    page_count = 0
    with fitz.open(pdf_path) as pdf:
        if pdf.page_count == 0:
            document.save(output_path)
            return 0

        first_page = pdf[0]
        section = document.sections[0]
        _set_section(section, float(first_page.rect.width), float(first_page.rect.height))

        for idx, page in enumerate(pdf):
            page_count += 1
            if idx > 0:
                prev = pdf[idx - 1]
                same_size = abs(prev.rect.width - page.rect.width) < 1 and abs(prev.rect.height - page.rect.height) < 1
                if same_size:
                    document.add_page_break()
                else:
                    section = document.add_section(WD_SECTION.NEW_PAGE)
                    _set_section(section, float(page.rect.width), float(page.rect.height))

            image_path = temp_dir / f"visual_{idx + 1:04d}.png"
            _render_page_image(page, image_path, dpi)

            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Pt(float(page.rect.width)), height=Pt(float(page.rect.height)))

    document.save(output_path)
    return page_count
