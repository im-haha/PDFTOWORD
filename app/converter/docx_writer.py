from __future__ import annotations

from pathlib import Path
from statistics import median

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.converter.schema import DocumentLayout, ImageData, PageLayout, ParagraphData, TableData


DEFAULT_MARGIN_PT = 36.0


def _set_section_size(section, page: PageLayout) -> None:
    section.page_width = Pt(page.width)
    section.page_height = Pt(page.height)
    section.left_margin = Pt(DEFAULT_MARGIN_PT)
    section.right_margin = Pt(DEFAULT_MARGIN_PT)
    section.top_margin = Pt(DEFAULT_MARGIN_PT)
    section.bottom_margin = Pt(DEFAULT_MARGIN_PT)


def _set_alignment(paragraph, align: str) -> None:
    if align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _append_run(paragraph, run_data) -> None:
    if run_data.text == "":
        return
    run = paragraph.add_run(run_data.text)
    run.font.name = run_data.font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), run_data.font_name)
    run.font.size = Pt(run_data.font_size)
    run.bold = run_data.bold
    run.italic = run_data.italic
    if run_data.color:
        r, g, b = run_data.color
        run.font.color.rgb = RGBColor(r, g, b)


def _write_paragraph(document: Document, data: ParagraphData) -> None:
    p = document.add_paragraph()
    fmt = p.paragraph_format
    fmt.left_indent = Pt(data.left_indent)
    fmt.right_indent = Pt(max(0.0, data.right_indent))
    fmt.first_line_indent = Pt(data.first_line_indent)
    fmt.space_before = Pt(data.space_before)
    fmt.space_after = Pt(data.space_after)
    fmt.line_spacing = Pt(max(1.0, data.line_spacing))

    _set_alignment(p, data.align)

    for run_data in data.runs:
        _append_run(p, run_data)


def _vertical_overlap_ratio(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> float:
    ay0, ay1 = a[1], a[3]
    by0, by1 = b[1], b[3]
    inter = max(0.0, min(ay1, by1) - max(ay0, by0))
    ah = max(1.0, ay1 - ay0)
    bh = max(1.0, by1 - by0)
    return inter / min(ah, bh)


def _same_row(a: ParagraphData, b: ParagraphData) -> bool:
    if abs(a.bbox[1] - b.bbox[1]) > 1.8:
        return False
    return _vertical_overlap_ratio(a.bbox, b.bbox) >= 0.6


def _group_paragraph_rows(paragraphs: list[ParagraphData]) -> list[list[ParagraphData]]:
    if not paragraphs:
        return []

    ordered = sorted(paragraphs, key=lambda p: (p.bbox[1], p.bbox[0]))
    rows: list[list[ParagraphData]] = []
    current: list[ParagraphData] = [ordered[0]]

    for para in ordered[1:]:
        if _same_row(current[0], para):
            current.append(para)
        else:
            rows.append(sorted(current, key=lambda p: p.bbox[0]))
            current = [para]

    if current:
        rows.append(sorted(current, key=lambda p: p.bbox[0]))

    return rows


def _write_row_paragraph(document: Document, row: list[ParagraphData]) -> None:
    if len(row) == 1:
        _write_paragraph(document, row[0])
        return

    first = row[0]
    p = document.add_paragraph()
    fmt = p.paragraph_format
    fmt.left_indent = Pt(max(0.0, first.bbox[0] - DEFAULT_MARGIN_PT))
    fmt.right_indent = Pt(0.0)
    fmt.first_line_indent = Pt(0.0)
    fmt.space_before = Pt(max(item.space_before for item in row))
    fmt.space_after = Pt(max(item.space_after for item in row))
    fmt.line_spacing = Pt(max(1.0, median(item.line_spacing for item in row)))
    _set_alignment(p, "left")

    for run_data in first.runs:
        _append_run(p, run_data)

    for item in row[1:]:
        tab_pos = max(1.0, item.bbox[0] - DEFAULT_MARGIN_PT)
        fmt.tab_stops.add_tab_stop(Pt(tab_pos))
        p.add_run("\t")
        for run_data in item.runs:
            _append_run(p, run_data)


def _write_table(document: Document, data: TableData) -> None:
    if not data.rows:
        return

    rows = len(data.rows)
    cols = max(len(row) for row in data.rows)
    if cols <= 0:
        return

    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"

    for r_idx, row in enumerate(data.rows):
        for c_idx in range(cols):
            text = row[c_idx] if c_idx < len(row) else ""
            table.cell(r_idx, c_idx).text = text


def _write_image(document: Document, page: PageLayout, image: ImageData) -> None:
    img_path = Path(image.image_path)
    if not img_path.exists():
        return

    p = document.add_paragraph()
    run = p.add_run()

    available_width = page.width - DEFAULT_MARGIN_PT * 2
    image_width = max(30.0, min(available_width, image.bbox[2] - image.bbox[0]))

    run.add_picture(str(img_path), width=Pt(image_width))


def _is_same_size(a: PageLayout, b: PageLayout) -> bool:
    return abs(a.width - b.width) < 1.0 and abs(a.height - b.height) < 1.0


def write_docx(layout: DocumentLayout, output_path: str) -> None:
    document = Document()

    if not layout.pages:
        document.save(output_path)
        return

    current_section = document.sections[0]
    _set_section_size(current_section, layout.pages[0])

    for page_index, page in enumerate(layout.pages):
        if page_index > 0:
            prev_page = layout.pages[page_index - 1]
            if _is_same_size(prev_page, page):
                document.add_page_break()
            else:
                current_section = document.add_section(WD_SECTION.NEW_PAGE)
                _set_section_size(current_section, page)

        elements: list[tuple[float, str, object]] = []
        paragraph_rows = _group_paragraph_rows(page.paragraphs)
        for row in paragraph_rows:
            elements.append((min(item.bbox[1] for item in row), "paragraph_row", row))
        for table in page.tables:
            elements.append((table.bbox[1], "table", table))
        for image in page.images:
            elements.append((image.bbox[1], "image", image))
        elements.sort(key=lambda x: x[0])

        for _, kind, payload in elements:
            if kind == "paragraph_row":
                _write_row_paragraph(document, payload)  # type: ignore[arg-type]
            elif kind == "table":
                _write_table(document, payload)  # type: ignore[arg-type]
            elif kind == "image":
                _write_image(document, page, payload)  # type: ignore[arg-type]

    document.save(output_path)
