#!/usr/bin/env python3
"""
简历格式化脚本
将 bbb.docx 按照标准求职简历排版要求重新排版，输出 bbb_formatted.docx

标准简历排版规范：
- 页面: A4 纸, 上下边距 2.0cm, 左右边距 2.2cm
- 姓名: 黑体, 22pt, 居中, 加粗
- 应聘职位: 宋体, 12pt, 居中
- 联系信息: 宋体, 10.5pt, 居中, 一行展示
- 一级标题 (如"工作经验"): 黑体, 14pt, 加粗, 带底部边框线
- 二级标题 (如公司/项目名): 宋体, 12pt, 加粗
- 正文内容: 宋体, 11pt, 行距 1.15 倍
- 英文/数字: Calibri
"""

import re
import copy
from docx import Document
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# ── 配置 ──────────────────────────────────────────────
INPUT_FILE = "bbb.docx"
OUTPUT_FILE = "bbb_formatted.docx"

# 页面
PAGE_TOP_MARGIN = Cm(2.0)
PAGE_BOTTOM_MARGIN = Cm(2.0)
PAGE_LEFT_MARGIN = Cm(2.2)
PAGE_RIGHT_MARGIN = Cm(2.2)

# 字体
CN_FONT_TITLE = "黑体"          # 姓名、一级标题
CN_FONT_BODY = "宋体"           # 正文、二级标题
EN_FONT = "Calibri"             # 英文/数字

# 字号
NAME_SIZE = Pt(22)              # 姓名
POSITION_SIZE = Pt(12)          # 应聘职位
CONTACT_SIZE = Pt(10.5)         # 联系方式
SECTION_HEADING_SIZE = Pt(14)   # 一级标题
SUB_HEADING_SIZE = Pt(12)       # 二级标题（时段/项目名/技术栈）
BODY_SIZE = Pt(11)              # 正文
SMALL_LABEL_SIZE = Pt(10.5)     # 小标签

# 行距
LINE_SPACING = 1.15

# 颜色
COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)
COLOR_HEADING = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_ACCENT = RGBColor(0x2B, 0x57, 0x9A)  # 一级标题用深蓝色


# ── 辅助函数 ──────────────────────────────────────────

def set_run_font(run, cn_font, en_font, size, bold=False, italic=False, color=COLOR_BLACK):
    """统一设置 run 的中英文字体、字号等。"""
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = en_font
    # 设置中文字体（东亚字体）
    r = run._element
    rPr = r.find(qn("w:rPr"))
    if rPr is None:
        rPr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
        r.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), cn_font)
    rFonts.set(qn("w:ascii"), en_font)
    rFonts.set(qn("w:hAnsi"), en_font)


def set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                         space_before=Pt(0), space_after=Pt(0),
                         line_spacing=LINE_SPACING, first_line_indent=None):
    """统一设置段落格式。"""
    fmt = paragraph.paragraph_format
    fmt.alignment = alignment
    fmt.space_before = space_before
    fmt.space_after = space_after
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line_spacing
    if first_line_indent is not None:
        fmt.first_line_indent = first_line_indent
    else:
        fmt.first_line_indent = Pt(0)
    # 清除段前分页
    fmt.page_break_before = False
    fmt.keep_with_next = False
    fmt.widow_control = True


def add_text_run(paragraph, text, cn_font, en_font, size,
                 bold=False, italic=False, color=COLOR_BLACK):
    """添加一个格式化的 run。"""
    run = paragraph.add_run(text)
    set_run_font(run, cn_font, en_font, size, bold, italic, color)
    return run


def add_section_heading(doc, title):
    """添加一级标题段落（如 "工作经验"），带底部分隔线。"""
    p = doc.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT,
                         space_before=Pt(14), space_after=Pt(6))
    add_text_run(p, title, CN_FONT_TITLE, EN_FONT, SECTION_HEADING_SIZE,
                 bold=True, color=COLOR_ACCENT)
    # 添加底部边框线
    pPr = p._element.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="2B579A"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    return p


def add_sub_heading(doc, text, bold=True, size=SUB_HEADING_SIZE):
    """添加二级标题（如时间段、职位、项目名）。"""
    p = doc.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT,
                         space_before=Pt(6), space_after=Pt(2))
    add_text_run(p, text, CN_FONT_BODY, EN_FONT, size, bold=bold)
    return p


def add_body_text(doc, text, indent=False, space_before=Pt(1), space_after=Pt(1)):
    """添加正文段落。"""
    p = doc.add_paragraph()
    first_indent = Cm(0.5) if indent else Pt(0)
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT,
                         space_before=space_before, space_after=space_after,
                         first_line_indent=first_indent)
    add_text_run(p, text, CN_FONT_BODY, EN_FONT, BODY_SIZE)
    return p


def add_tech_stack_line(doc, text):
    """添加技术栈行（加粗、略小字号）。"""
    p = doc.add_paragraph()
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT,
                         space_before=Pt(2), space_after=Pt(4))
    add_text_run(p, text, CN_FONT_BODY, EN_FONT, SMALL_LABEL_SIZE, bold=True,
                 color=RGBColor(0x33, 0x33, 0x33))
    return p


# ── 内容提取 ──────────────────────────────────────────

def extract_content(doc):
    """从原始 docx 提取结构化内容。"""
    paragraphs = doc.paragraphs
    texts = [p.text.strip() for p in paragraphs]

    content = {
        "position": "",    # 应聘职位
        "name": "",
        "gender_age": "",
        "contact": "",
        "work_experience": [],   # list of {period, title, tech, duties:[]}
        "skills": [],            # list of str
        "projects": [],          # list of {name_tech, desc, duties:[]}
        "self_eval": "",
    }

    # --- 应聘职位 ---
    for t in texts:
        if t.startswith("应聘"):
            content["position"] = t.replace("应聘：", "").replace("应聘:", "").strip()
            break

    # --- 姓名 ---
    content["name"] = texts[1] if len(texts) > 1 else ""

    # --- 性别年龄 ---
    for t in texts:
        if "性别" in t and "年龄" in t:
            content["gender_age"] = t
            break

    # --- 联系方式 ---
    for t in texts:
        if "电话" in t and "邮箱" in t:
            content["contact"] = t
            break

    # --- 定位各节的起止索引 ---
    section_starts = {}
    section_order = []
    for i, t in enumerate(texts):
        if t.startswith("一、"):
            section_starts["work"] = i
            section_order.append("work")
        elif t.startswith("二、"):
            section_starts["skill"] = i
            section_order.append("skill")
        elif t.startswith("三、"):
            section_starts["project"] = i
            section_order.append("project")
        elif t.startswith("四、"):
            section_starts["eval"] = i
            section_order.append("eval")

    def section_range(key):
        idx = section_order.index(key)
        start = section_starts[key]
        end = section_starts[section_order[idx + 1]] if idx + 1 < len(section_order) else len(texts)
        return start, end

    # --- 工作经验 ---
    if "work" in section_starts:
        s, e = section_range("work")
        _parse_work_experience(texts[s + 1:e], content)

    # --- 专业技能 ---
    if "skill" in section_starts:
        s, e = section_range("skill")
        for t in texts[s + 1:e]:
            t = t.strip()
            if t:
                content["skills"].append(t)

    # --- 项目经验 ---
    if "project" in section_starts:
        s, e = section_range("project")
        _parse_projects(texts[s + 1:e], content)

    # --- 自我评价 ---
    if "eval" in section_starts:
        s, e = section_range("eval")
        eval_text = ""
        for t in texts[s + 1:e]:
            t = t.strip()
            if t:
                eval_text += t
        content["self_eval"] = eval_text

    return content


def _parse_work_experience(lines, content):
    """解析工作经验段落。"""
    current = None
    for line in lines:
        line = line.strip()
        if not line:
            continue
        # 时间段 (e.g. "2022.04-今" or "2019.06-2022.04")
        if re.match(r"^\d{4}\.\d{2}\s*[-–—]\s*(\d{4}\.\d{2}|今|至今)", line):
            if current:
                content["work_experience"].append(current)
            current = {"period": line, "title": "", "tech": "", "duties": []}
        elif current is not None:
            if not current["title"]:
                current["title"] = line
            elif not current["tech"] and any(kw in line for kw in ["Vue", "React", "Node", "Vite", "Pinia", "Vuex"]):
                current["tech"] = line
            else:
                # 职责条目
                current["duties"].append(line)
    if current:
        content["work_experience"].append(current)


def _parse_projects(lines, content):
    """解析项目经验段落。"""
    current = None
    in_desc = False
    in_duties = False

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # 项目标题行：以数字+、开头，且包含技术栈关键词
        if re.match(r"^\d+[、.]", line) and "技术栈" in line:
            if current:
                content["projects"].append(current)
            current = {"name_tech": line, "desc": "", "duties": []}
            in_desc = False
            in_duties = False
            continue

        if current is None:
            continue

        if line.startswith("项目描述"):
            in_desc = True
            in_duties = False
            # 项目描述可能在同一行
            rest = line.replace("项目描述：", "").replace("项目描述:", "").strip()
            if rest:
                current["desc"] += rest
            continue

        if line.startswith("工作内容与成果"):
            in_desc = False
            in_duties = True
            continue

        if in_duties:
            current["duties"].append(line)
        elif in_desc:
            current["desc"] += line
        # 其它内容也归入描述
        elif not in_desc and not in_duties:
            current["desc"] += line

    if current:
        content["projects"].append(current)


# ── 合并断行的技能条目 ──────────────────────────────────

def merge_skill_items(skills):
    """
    将因 PDF 断行而拆开的技能条目重新合并。
    以数字+、或数字+. 开头的视为新条目，否则追加到上一条。
    """
    merged = []
    for s in skills:
        if re.match(r"^\d+[、.]", s):
            merged.append(s)
        elif merged:
            merged[-1] += s
        else:
            merged.append(s)
    return merged


def merge_duty_items(duties):
    """合并断行的工作职责条目。"""
    merged = []
    for d in duties:
        if re.match(r"^\d+[\.\、]", d):
            merged.append(d)
        elif merged:
            merged[-1] += d
        else:
            merged.append(d)
    return merged


# ── 构建新文档 ──────────────────────────────────────────

def build_resume(content):
    """根据提取的内容构建标准格式简历。"""
    doc = Document()

    # ── 页面设置 ──
    section = doc.sections[0]
    section.page_width = Cm(21.0)   # A4
    section.page_height = Cm(29.7)  # A4
    section.top_margin = PAGE_TOP_MARGIN
    section.bottom_margin = PAGE_BOTTOM_MARGIN
    section.left_margin = PAGE_LEFT_MARGIN
    section.right_margin = PAGE_RIGHT_MARGIN

    # ── 设置默认样式 ──
    style = doc.styles["Normal"]
    style.font.name = EN_FONT
    style.font.size = BODY_SIZE
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT_BODY)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = LINE_SPACING

    # ============================================================
    # 1. 姓名
    # ============================================================
    p_name = doc.add_paragraph()
    set_paragraph_format(p_name, WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=Pt(0), space_after=Pt(4))
    add_text_run(p_name, content["name"], CN_FONT_TITLE, EN_FONT,
                 NAME_SIZE, bold=True, color=COLOR_HEADING)

    # ============================================================
    # 2. 应聘职位
    # ============================================================
    p_pos = doc.add_paragraph()
    set_paragraph_format(p_pos, WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=Pt(2), space_after=Pt(4))
    add_text_run(p_pos, f"应聘：{content['position']}", CN_FONT_BODY, EN_FONT,
                 POSITION_SIZE, color=RGBColor(0x33, 0x33, 0x33))

    # ============================================================
    # 3. 联系信息（合并为一行，居中）
    # ============================================================
    contact_parts = []
    ga = content["gender_age"]
    if ga:
        # 提取性别和年龄
        m_gender = re.search(r"性别[：:]\s*(\S+)", ga)
        m_age = re.search(r"年龄[：:]\s*(\S+)", ga)
        if m_gender:
            contact_parts.append(f"性别：{m_gender.group(1)}")
        if m_age:
            contact_parts.append(f"年龄：{m_age.group(1)}")

    ct = content["contact"]
    if ct:
        m_phone = re.search(r"电话[：:]\s*(\S+)", ct)
        m_email = re.search(r"邮箱[：:]\s*(\S+)", ct)
        if m_phone:
            contact_parts.append(f"电话：{m_phone.group(1)}")
        if m_email:
            contact_parts.append(f"邮箱：{m_email.group(1)}")

    if contact_parts:
        p_contact = doc.add_paragraph()
        set_paragraph_format(p_contact, WD_ALIGN_PARAGRAPH.CENTER,
                             space_before=Pt(0), space_after=Pt(10))
        separator = "    |    "
        add_text_run(p_contact, separator.join(contact_parts),
                     CN_FONT_BODY, EN_FONT, CONTACT_SIZE,
                     color=RGBColor(0x55, 0x55, 0x55))

    # ============================================================
    # 4. 工作经验
    # ============================================================
    if content["work_experience"]:
        add_section_heading(doc, "工作经验")
        for exp in content["work_experience"]:
            # 时间段 + 职位标题同行
            p = doc.add_paragraph()
            set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT,
                                 space_before=Pt(8), space_after=Pt(2))
            add_text_run(p, exp["period"], CN_FONT_BODY, EN_FONT,
                         SUB_HEADING_SIZE, bold=True)
            add_text_run(p, "    ", CN_FONT_BODY, EN_FONT, SUB_HEADING_SIZE)
            add_text_run(p, exp["title"], CN_FONT_BODY, EN_FONT,
                         SUB_HEADING_SIZE, bold=True)

            # 技术栈
            if exp["tech"]:
                add_tech_stack_line(doc, f"技术栈：{exp['tech']}")

            # 工作职责
            duties = merge_duty_items(exp["duties"])
            for duty in duties:
                add_body_text(doc, duty)

    # ============================================================
    # 5. 专业技能
    # ============================================================
    if content["skills"]:
        add_section_heading(doc, "专业技能")
        merged_skills = merge_skill_items(content["skills"])
        for skill in merged_skills:
            add_body_text(doc, skill, space_before=Pt(2), space_after=Pt(2))

    # ============================================================
    # 6. 项目经验
    # ============================================================
    if content["projects"]:
        add_section_heading(doc, "项目经验")
        for proj in content["projects"]:
            # 项目标题 + 技术栈
            name_tech = proj["name_tech"]
            # 拆分项目名和技术栈
            tech_match = re.search(r"技术栈[：:](.+)", name_tech)
            proj_name_match = re.match(r"(\d+[、.].+?)技术栈", name_tech)

            p = doc.add_paragraph()
            set_paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT,
                                 space_before=Pt(10), space_after=Pt(2))

            proj_name = proj_name_match.group(1).strip() if proj_name_match else name_tech
            add_text_run(p, proj_name, CN_FONT_BODY, EN_FONT,
                         SUB_HEADING_SIZE, bold=True)

            if tech_match:
                tech_text = tech_match.group(1).strip()
                add_tech_stack_line(doc, f"技术栈：{tech_text}")

            # 项目描述
            if proj["desc"]:
                p_desc_label = doc.add_paragraph()
                set_paragraph_format(p_desc_label, WD_ALIGN_PARAGRAPH.LEFT,
                                     space_before=Pt(4), space_after=Pt(1))
                add_text_run(p_desc_label, "项目描述：", CN_FONT_BODY, EN_FONT,
                             BODY_SIZE, bold=True)

                add_body_text(doc, proj["desc"], indent=True,
                              space_before=Pt(1), space_after=Pt(4))

            # 工作内容
            duties = merge_duty_items(proj["duties"])
            if duties:
                p_duty_label = doc.add_paragraph()
                set_paragraph_format(p_duty_label, WD_ALIGN_PARAGRAPH.LEFT,
                                     space_before=Pt(4), space_after=Pt(1))
                add_text_run(p_duty_label, "工作内容与成果：", CN_FONT_BODY, EN_FONT,
                             BODY_SIZE, bold=True)

                for duty in duties:
                    add_body_text(doc, duty)

    # ============================================================
    # 7. 自我评价
    # ============================================================
    if content["self_eval"]:
        add_section_heading(doc, "自我评价")
        add_body_text(doc, content["self_eval"], indent=True,
                      space_before=Pt(4), space_after=Pt(4))

    return doc


# ── 主入口 ──────────────────────────────────────────

def main():
    print(f"读取源文件：{INPUT_FILE}")
    src = Document(INPUT_FILE)

    print("提取简历内容...")
    content = extract_content(src)

    print(f"  姓名: {content['name']}")
    print(f"  应聘: {content['position']}")
    print(f"  工作经验: {len(content['work_experience'])} 段")
    print(f"  专业技能: {len(content['skills'])} 条")
    print(f"  项目经验: {len(content['projects'])} 个")

    print("构建标准格式简历...")
    doc = build_resume(content)

    doc.save(OUTPUT_FILE)
    print(f"\n✅ 格式化完成，已输出: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
